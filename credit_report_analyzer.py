# -*- coding: utf-8 -*-
"""授信报告 OCR、财务交叉分析与综合风险输出。"""
import json
import hashlib
import logging
from pathlib import Path
from time import monotonic

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from cache_manager import CacheManager
from llm_client import extract_json
from report_markdown import write_markdown_report
from scanner import find_financial_dirs, list_images
from summary_manager import safe_filename

logger = logging.getLogger(__name__)

CREDIT_EXTRACTION_PROMPT = """你是授信报告信息抽取助手。请从图片中逐项提取可见事实，不做推断，不补造缺失内容。
只返回 JSON 对象，日期尽量保留原文，金额保留数值与单位。字段无法确认时填 null。
字段：report_date（报告日期/基准日）、project_total_scale（项目总规模）、
project_investment_amount（项目总投入/总投资）、invested_amount（已投入/已完成投资）、
remaining_investment（尚未投入/剩余投资）、project_status（项目状态）、
construction_progress（建设进度）、unsold_units（待销售楼盘/房源数量）、
unsold_ratio（待销售数量占比，保留百分比原文）、tail_end（是否属于尾盘）、
follow_on_funding_sufficiency（后续资金是否充足及依据）、funding_sources（后续资金来源）、
key_conditions（授信额度、提款条件、期限、担保等关键条件）、
risks_disclosed（授信报告明确披露的风险）、evidence（关键事实对应的原文短摘录）。
图片序号：{image_name}
"""

CROSS_ANALYSIS_PROMPT = """你是房地产项目授信与财务交叉审查分析师。请只依据给定的授信报告抽取结果和财务报表数据分析【{project_name}】。

时间规则：每个授信事实按其报告日期/基准日与财报年度对齐。没有具体日期时标为时间不明；不得把报告时点之后的信息倒用于早期财报。发现年份/统计口径不一致时明确披露。

重点核查并给出事实、来源与判断：
1. 项目总规模、项目总投入、已投入、剩余投入及金额勾稽；如数字口径不同，不要强行相减。
2. 项目状态、建设进度、待销售楼盘/房源数量及占比、是否尾盘。
3. 后续资金是否充足，结合项目剩余投资、企业货币资金、负债、有息债务、经营现金流及授信提款条件判断；区分已批授信、可提款资金和实际到账资金。
4. 授信报告已披露的风险，以及授信报告未反映但财务报表显示的风险。逐项列出未披露/弱化的财务风险及财报年份、对应科目或数据；无法确认时写“待核实”，不得断言隐瞒。
5. 给出综合风险等级（高/中/低）、0-100 分（越高风险越大）、简要结论、主要风险、建议和资料局限。

只返回符合以下字段的 JSON 对象；未知数值填 null，数组无内容时填空数组：
{{
  "project_name": "{project_name}",
  "report_date": null,
  "project_total_scale": null,
  "project_investment_amount": null,
  "invested_amount": null,
  "remaining_investment": null,
  "project_status": null,
  "construction_progress": null,
  "unsold_units": null,
  "unsold_ratio": null,
  "tail_end": null,
  "follow_on_funding_sufficiency": null,
  "funding_assessment": "",
  "credit_disclosed_risks": [],
  "financial_risks_not_reflected": [],
  "cross_checks": [],
  "risk_level": "中",
  "risk_score": 50,
  "summary": "",
  "key_risks": [],
  "suggestions": [],
  "limitations": []
}}

授信报告抽取结果：
{credit_facts}

已生成的项目财务报表明细：
{financial_data}
"""


def extract_credit_facts(llm, cache: CacheManager, credit_dirs, image_exts) -> list:
    images = []
    for directory in credit_dirs:
        images.extend(list_images(directory, image_exts))
    images = sorted(set(images), key=lambda path: str(path).lower())
    facts = []
    cache_hits = 0
    failed = 0
    logger.info("授信报告图片：共 %d 张，缓存命中检查开始", len(images))

    for index, image in enumerate(images, 1):
        image_digest = hashlib.sha256(image.read_bytes()).hexdigest()
        cache_key = f"credit_ocr::{image.resolve()}::{image_digest}"
        cache_label = f"{image.parents[1].name}_授信报告OCR_{image.stem}"
        record = cache.get(cache_key, label=cache_label)
        if record and isinstance(record.get("parsed"), dict) and "_parse_error" not in record["parsed"]:
            facts.append(record["parsed"])
            cache_hits += 1
            logger.info("  授信图片 %d/%d：缓存命中 [%s]", index, len(images), image.name)
            continue

        logger.info("  授信图片 %d/%d：调用模型抽取 [%s]", index, len(images), image.name)
        prompt = CREDIT_EXTRACTION_PROMPT.format(image_name=image.name)
        raw = llm.chat_with_images([str(image)], prompt)
        try:
            finish_reason = getattr(llm, "last_finish_reason", None)
            if finish_reason not in (None, "stop"):
                raise ValueError(f"模型结束原因：{finish_reason}")
            parsed = extract_json(raw)
            if not isinstance(parsed, dict):
                raise ValueError("授信报告抽取结果不是 JSON 对象")
        except Exception as exc:  # noqa: BLE001
            logger.error("授信报告图片 [%s] 抽取失败：%s", image, exc)
            failed += 1
            continue

        cache.set(
            cache_key,
            {"images": [str(image)], "raw": raw, "parsed": parsed},
            label=cache_label,
        )
        facts.append(parsed)
        logger.info("  授信图片 %d/%d：抽取成功", index, len(images))

    logger.info(
        "授信报告抽取完成：有效 %d，缓存命中 %d，失败 %d",
        len(facts), cache_hits, failed,
    )
    return facts


def load_financial_workbook(path, max_rows: int = 5000) -> list:
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise FileNotFoundError(f"缺少已生成的项目财务报表：{workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook["明细"]
        rows = []
        for values in worksheet.iter_rows(min_row=2, values_only=True):
            if not any(value is not None for value in values):
                continue
            rows.append({
                "project": values[0],
                "company": values[1],
                "year": values[2],
                "statement_type": values[3],
                "unit": values[4],
                "item": values[5],
                "column": values[6],
                "value": values[7],
            })
            if len(rows) >= max_rows:
                logger.warning("财务明细超过 %d 行，交叉分析仅使用前 %d 行", max_rows, max_rows)
                break
        return rows
    finally:
        workbook.close()


def analyze_comprehensive_risk(llm, project_name: str, credit_facts: list,
                               financial_rows: list) -> dict:
    prompt = CROSS_ANALYSIS_PROMPT.format(
        project_name=project_name,
        credit_facts=json.dumps(credit_facts, ensure_ascii=False, indent=2),
        financial_data=json.dumps(financial_rows, ensure_ascii=False, indent=2),
    )
    logger.info(
        "综合分析请求：授信事实 %d 项，财务明细 %d 行",
        len(credit_facts), len(financial_rows),
    )
    raw = llm.chat([{"role": "user", "content": prompt}], temperature=0.2)
    result = extract_json(raw)
    if not isinstance(result, dict):
        raise ValueError("综合风险分析结果不是 JSON 对象")
    return result


def write_comprehensive_summary(records: list, json_path, excel_path) -> None:
    json_path = Path(json_path)
    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(
        json.dumps({"projects": records}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "综合风险汇总"
    headers = [
        "项目名称", "授信报告日期", "项目总规模", "项目总投入", "已投入金额",
        "尚未投入金额", "项目状态", "建设进度", "待销售数量", "待销售占比",
        "是否尾盘", "后续资金判断", "综合风险等级", "风险评分",
        "授信未反映的财务风险", "主要风险", "建议", "项目报告",
    ]
    worksheet.append(headers)
    for record in records:
        risks = record.get("financial_risks_not_reflected") or []
        key_risks = record.get("key_risks") or []
        suggestions = record.get("suggestions") or []
        worksheet.append([
            record.get("project_name"), record.get("report_date"),
            record.get("project_total_scale"), record.get("project_investment_amount"),
            record.get("invested_amount"), record.get("remaining_investment"),
            record.get("project_status"), record.get("construction_progress"),
            record.get("unsold_units"), record.get("unsold_ratio"),
            record.get("tail_end"), record.get("follow_on_funding_sufficiency"),
            record.get("risk_level"), record.get("risk_score"),
            "\n".join(map(str, risks)), "\n".join(map(str, key_risks)),
            "\n".join(map(str, suggestions)), record.get("report_path"),
        ])

    widths = [24, 16, 16, 16, 16, 16, 18, 16, 14, 14, 12, 28, 12, 12, 48, 42, 42, 48]
    for index, width in enumerate(widths, 1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    excel_path = Path(excel_path)
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(excel_path)


def run_credit_report_analysis(llm, cache, projects, project_output_root,
                               image_exts, dir_keywords, max_depth,
                               summary_json_path, summary_excel_path,
                               project_excel_name, report_name,
                               force_rerun=False) -> list:
    started = monotonic()
    projects = list(projects)
    records = []
    counts = {"analyzed": 0, "reused": 0, "skipped": 0, "failed": 0}
    logger.info("开始授信报告综合分析：扫描到 %d 个项目", len(projects))
    for project_index, project in enumerate(projects, 1):
        project_name = project.name
        logger.info("[%d/%d] 项目 [%s]：检查授信目录", project_index, len(projects), project_name)
        credit_dirs = find_financial_dirs(project, dir_keywords, max_depth)
        if not credit_dirs:
            logger.info("项目 [%s] 未找到授信报告目录，跳过", project_name)
            counts["skipped"] += 1
            continue
        logger.info("项目 [%s]：找到 %d 个授信目录", project_name, len(credit_dirs))

        project_dir = Path(project_output_root) / safe_filename(project_name)
        financial_path = project_dir / project_excel_name
        if not financial_path.is_file():
            logger.warning("项目 [%s] 尚无财务报表 Excel，无法做综合分析", project_name)
            counts["skipped"] += 1
            continue
        logger.info("项目 [%s]：读取财务报表 [%s]", project_name, financial_path.name)

        fingerprint_payload = []
        stat = financial_path.stat()
        fingerprint_payload.append((str(financial_path.resolve()), stat.st_size, stat.st_mtime_ns))
        for directory in credit_dirs:
            for image in list_images(directory, image_exts):
                stat = image.stat()
                fingerprint_payload.append((str(image.resolve()), stat.st_size, stat.st_mtime_ns))
        fingerprint = hashlib.sha256(
            json.dumps(fingerprint_payload, ensure_ascii=False).encode("utf-8")
        ).hexdigest()
        report_path = project_dir / report_name
        if not force_rerun and report_path.is_file():
            try:
                previous = json.loads(report_path.read_text(encoding="utf-8"))
                if previous.get("fingerprint") == fingerprint:
                    write_markdown_report(
                        previous,
                        report_path.with_suffix(".md"),
                        f"{project_name}综合风险报告",
                    )
                    records.append(previous["risk_result"] | {"report_path": str(report_path)})
                    logger.info("项目 [%s] 输入未变化，复用综合风险报告", project_name)
                    counts["reused"] += 1
                    continue
            except (OSError, ValueError, KeyError, TypeError):
                pass

        try:
            credit_facts = extract_credit_facts(llm, cache, credit_dirs, image_exts)
        except Exception as exc:  # noqa: BLE001
            logger.exception("项目 [%s] 授信报告识别失败：%s", project_name, exc)
            counts["failed"] += 1
            continue
        if not credit_facts:
            logger.warning("项目 [%s] 没有可用的授信报告抽取结果，跳过", project_name)
            counts["skipped"] += 1
            continue

        try:
            financial_rows = load_financial_workbook(financial_path)
            logger.info("项目 [%s]：开始授信与财务交叉分析", project_name)
            result = analyze_comprehensive_risk(llm, project_name, credit_facts, financial_rows)
        except Exception as exc:  # noqa: BLE001
            logger.exception("项目 [%s] 综合风险分析失败：%s", project_name, exc)
            counts["failed"] += 1
            continue

        project_dir.mkdir(parents=True, exist_ok=True)
        report = {
            "project_name": project_name,
            "fingerprint": fingerprint,
            "credit_source_dirs": [str(path) for path in credit_dirs],
            "credit_facts": credit_facts,
            "financial_source": str(financial_path),
            "risk_result": result,
        }
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        markdown_path = write_markdown_report(
            report,
            report_path.with_suffix(".md"),
            f"{project_name}综合风险报告",
        )
        records.append(result | {"report_path": str(report_path)})
        counts["analyzed"] += 1
        logger.info(
            "项目 [%s]：综合分析完成，风险=%s，评分=%s，报告=%s",
            project_name, result.get("risk_level", "未知"),
            result.get("risk_score", "未知"), markdown_path,
        )

    write_comprehensive_summary(records, summary_json_path, summary_excel_path)
    logger.info("综合风险 JSON 汇总：%s", summary_json_path)
    logger.info("综合风险 Excel 汇总：%s", summary_excel_path)
    logger.info(
        "授信综合分析结束：分析 %d，复用 %d，跳过 %d，失败 %d，用时 %.1f 秒",
        counts["analyzed"], counts["reused"], counts["skipped"],
        counts["failed"], monotonic() - started,
    )
    return records