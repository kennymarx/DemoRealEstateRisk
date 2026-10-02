# -*- coding: utf-8 -*-
"""风险信息汇总表管理：备份 + 断点续作"""
import hashlib
import json
import logging
import re
import threading
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

SUMMARY_VERSION = 1

LEVEL_FILLS = {
    "高": PatternFill("solid", fgColor="FFC7CE"),
    "中": PatternFill("solid", fgColor="FFEB9C"),
    "低": PatternFill("solid", fgColor="C6EFCE"),
    "未知": PatternFill("solid", fgColor="D9D9D9"),
}
LEVEL_FONTS = {
    "高": Font(color="9C0006", bold=True),
    "中": Font(color="9C6500", bold=True),
    "低": Font(color="006100", bold=True),
    "未知": Font(color="595959", bold=True),
}


def safe_filename(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "_", str(name)).strip() or "unnamed"


class SummaryManager:
    def __init__(self, json_path, excel_path=None,
                 backup_dir=None, backup_keep: int = 50):
        self.json_path = Path(json_path)
        self.excel_path = Path(excel_path) if excel_path else self.json_path.with_suffix(".xlsx")
        self.backup_dir = Path(backup_dir) if backup_dir else self.json_path.parent / "history"
        self.backup_keep = int(backup_keep)
        self._lock = threading.RLock()
        self.data = self._load()

    # ---------- 备份 ----------
    def _backup_current(self) -> None:
        if self.backup_keep <= 0 or not self.json_path.exists():
            return
        try:
            raw = self.json_path.read_bytes()
        except OSError as e:
            logger.warning("读取待备份文件失败：%s", e)
            return

        try:
            self.backup_dir.mkdir(parents=True, exist_ok=True)
            existing = sorted(self.backup_dir.glob("汇总表_*.json"))
        except OSError:
            existing = []

        if existing:
            try:
                if existing[-1].read_bytes() == raw:
                    return
            except OSError:
                pass

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        dst = self.backup_dir / f"汇总表_{ts}.json"
        i = 1
        while dst.exists():
            dst = self.backup_dir / f"汇总表_{ts}_{i}.json"
            i += 1

        try:
            dst.write_bytes(raw)
        except OSError as e:
            logger.warning("备份失败：%s", e)
            return

        self._cleanup_backups()

    def _cleanup_backups(self) -> None:
        if self.backup_keep <= 0:
            return
        try:
            files = sorted(self.backup_dir.glob("汇总表_*.json"))
        except OSError:
            return
        for old in files[:-self.backup_keep]:
            try:
                old.unlink()
            except OSError:
                pass

    def list_backups(self) -> list:
        try:
            return sorted(self.backup_dir.glob("汇总表_*.json"))
        except OSError:
            return []

    def restore_from_backup(self, backup_file) -> None:
        p = Path(backup_file)
        if not p.exists():
            raise FileNotFoundError(p)
        self._backup_current()
        self.data = json.loads(p.read_text("utf-8"))
        self.save()
        logger.warning("已从备份恢复：%s", p)

    # ---------- 读写 ----------
    def _load(self) -> dict:
        if not self.json_path.exists():
            return {"version": SUMMARY_VERSION, "updated_at": None, "projects": {}}
        try:
            d = json.loads(self.json_path.read_text("utf-8"))
            if not isinstance(d, dict):
                raise ValueError("根节点不是对象")
            d.setdefault("version", SUMMARY_VERSION)
            d.setdefault("projects", {})
            return d
        except Exception as e:  # noqa: BLE001
            logger.warning("汇总表读取失败（将重建）：%s", e)
            return {"version": SUMMARY_VERSION, "updated_at": None, "projects": {}}

    def save(self) -> None:
        with self._lock:
            self._backup_current()
            self.data["updated_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            self.json_path.parent.mkdir(parents=True, exist_ok=True)
            tmp = self.json_path.with_suffix(".json.tmp")
            tmp.write_text(
                json.dumps(self.data, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            tmp.replace(self.json_path)

    # ---------- 断点续作 ----------
    def is_done(self, project_name: str, fingerprint: str) -> bool:
        rec = self.data["projects"].get(project_name)
        if not rec:
            return False
        if rec.get("status") != "success":
            return False
        return rec.get("fingerprint") == fingerprint

    def get(self, project_name: str) -> Optional[dict]:
        return self.data["projects"].get(project_name)

    def all_projects(self) -> dict:
        return dict(self.data["projects"])

    # ---------- 更新 ----------
    def upsert(self, project_name: str, **kwargs) -> None:
        with self._lock:
            rec = dict(self.data["projects"].get(project_name) or {})
            rec["project_name"] = project_name
            rec.update(kwargs)
            rec["processed_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            self.data["projects"][project_name] = rec
        self.save()

    def mark_failed(self, project_name: str, fingerprint: str, error: str,
                    attempt: int = 1) -> None:
        prev = self.data["projects"].get(project_name) or {}
        self.upsert(
            project_name,
            fingerprint=fingerprint,
            status="failed",
            attempt=int(attempt),
            error=str(error)[:500],
            last_success_at=prev.get("processed_at") if prev.get("status") == "success" else prev.get("last_success_at"),
        )

    # ---------- 导出 Excel ----------
    def export_excel(self) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "风险信息汇总"

        headers = [
            "项目名称", "开发商", "报表年份", "报表数量",
            "风险等级", "风险评分",
            "关键风险点", "分析摘要",
            "状态", "尝试次数", "处理时间", "报告文件", "备注",
        ]
        ws.append(headers)

        def sort_key(item):
            r = item[1]
            try:
                score = float(r.get("risk_score") or 0)
            except (TypeError, ValueError):
                score = 0
            return (-score, item[0])

        for name, r in sorted(self.data["projects"].items(), key=sort_key):
            devs = r.get("developers")
            devs_txt = "、".join(str(x) for x in devs) if isinstance(devs, list) else (devs or "")
            years = r.get("report_years")
            years_txt = "、".join(str(x) for x in sorted(set(years))) \
                if isinstance(years, list) else (years or "")
            krs = r.get("key_risks") or []
            krs_txt = "\n".join(f"{i + 1}. {k}" for i, k in enumerate(krs)) \
                if isinstance(krs, list) and krs else ""

            ws.append([
                name, devs_txt, years_txt, r.get("statement_count") or 0,
                r.get("risk_level") or "", r.get("risk_score"),
                krs_txt, r.get("summary") or "",
                r.get("status") or "", r.get("attempt") or 1,
                r.get("processed_at") or "",
                r.get("report_path") or "", r.get("error") or "",
            ])

        widths = [24, 26, 14, 10, 10, 10, 48, 40, 10, 10, 20, 32, 28]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_idx in range(2, ws.max_row + 1):
            lv_cell = ws.cell(row=row_idx, column=5)
            lv = str(lv_cell.value or "").strip()
            if lv in LEVEL_FILLS:
                lv_cell.fill = LEVEL_FILLS[lv]
                lv_cell.font = LEVEL_FONTS[lv]
                lv_cell.alignment = Alignment(horizontal="center", vertical="center")

            if str(ws.cell(row=row_idx, column=9).value or "") == "failed":
                for c in range(1, len(headers) + 1):
                    if c != 5:
                        ws.cell(row=row_idx, column=c).fill = PatternFill(
                            "solid", fgColor="F2F2F2"
                        )

            ws.cell(row=row_idx, column=6).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=10).alignment = Alignment(horizontal="center")
            for c in (7, 8, 12, 13):
                ws.cell(row=row_idx, column=c).alignment = Alignment(
                    wrap_text=True, vertical="top"
                )

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        self.excel_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(self.excel_path)
        return self.excel_path

    # ---------- 简报 ----------
    def print_summary(self) -> None:
        projects = self.data["projects"]
        if not projects:
            print("（汇总表为空）")
            return
        levels = {"高": 0, "中": 0, "低": 0, "未知": 0}
        failed = 0
        for r in projects.values():
            lv = str(r.get("risk_level") or "未知").strip()
            levels[lv] = levels.get(lv, 0) + 1
            if r.get("status") == "failed":
                failed += 1
        print(
            f"共 {len(projects)} 个项目 | "
            f"🔴 高风险 {levels.get('高', 0)} | "
            f"🟡 中风险 {levels.get('中', 0)} | "
            f"🟢 低风险 {levels.get('低', 0)} | "
            f"⚪ 未知 {levels.get('未知', 0)} | "
            f"❌ 失败 {failed}"
        )
        print(f"   JSON     : {self.json_path}")
        print(f"   Excel    : {self.excel_path}")
        if self.backup_keep > 0:
            backups = self.list_backups()
            print(f"   备份目录 : {self.backup_dir}  (共 {len(backups)} 个，"
                  f"保留最近 {self.backup_keep} 个)")


# ---------- 项目指纹 ----------
def project_fingerprint(project_dir, fs_dirs, image_exts) -> str:
    from scanner import list_images
    h = hashlib.md5()
    h.update(str(Path(project_dir).resolve()).encode("utf-8"))
    for fs_dir in sorted(Path(d) for d in fs_dirs):
        h.update(b"|DIR|")
        h.update(str(fs_dir.resolve()).encode("utf-8"))
        for img in list_images(fs_dir, image_exts):
            try:
                st = img.stat()
                h.update(f"|{img.resolve()}|{st.st_size}|{int(st.st_mtime)}".encode("utf-8"))
            except OSError:
                h.update(f"|{img}".encode("utf-8"))
    return h.hexdigest()